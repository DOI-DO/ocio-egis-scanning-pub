#!/usr/bin/env python3
r"""
PORTAL_ROLES.py
==============================================================================
Generate a privilege matrix for every custom role in Portal or ArcGIS Online.

Output: creates a timestamped copy of ROLES_TEMPLATE.xlsx and updates its
"MATRIX" sheet with a permission matrix:
  * Columns starting at D: role names, ordered by permission count.
  * Column C: existing privilege labels; this script never changes them.
  * Matrix cells: a checkmark when a role has the privilege, otherwise blank.

The privilege-label-to-REST-code mapping is embedded in PRIVILEGE_DATA. Some
ArcGIS releases use different privilege codes for the same displayed label;
such codes are treated as aliases rather than silently overwriting one another.

All direct REST traffic and the ArcGIS SDK connection use the same explicit TLS
verification strategy selected with --tls-mode. REST failures invalidate the
run rather than allowing a partial role matrix to be reported as complete.

Requires --config <org_toml> and --target {portal,agol}; settings are resolved
from the selected target section plus [common].

Examples:
    python PORTAL_ROLES.py --config ORG.config.toml --target portal
    python PORTAL_ROLES.py --config ORG.config.toml --target agol
    python PORTAL_ROLES.py --config ORG.config.toml --target portal \
        --tls-mode os-trust-store
    python PORTAL_ROLES.py --config ORG.config.toml --target portal \
        --tls-mode custom-ca-bundle --ca-bundle-path C:\certs\internal-ca.pem
"""

from __future__ import annotations

import argparse
import importlib.util
import os
import sys
import tempfile
from pathlib import Path
from urllib.parse import unquote, urlsplit, urlunsplit


# Locate the project root before importing package code. Production scripts may
# live either directly in the project root or in a one-level child directory.
_SCRIPT_DIR = os.path.dirname(os.path.realpath(__file__))
_PARENT_DIR = os.path.dirname(_SCRIPT_DIR)
if os.path.isdir(os.path.join(_PARENT_DIR, "UTILITIES")):
    _PROJECT_ROOT = _PARENT_DIR
elif os.path.isdir(os.path.join(_SCRIPT_DIR, "UTILITIES")):
    _PROJECT_ROOT = _SCRIPT_DIR
else:
    _PROJECT_ROOT = _PARENT_DIR


def _path_key(value):
    """Return a normalized real path for safe sys.path comparisons."""
    try:
        return os.path.normcase(os.path.realpath(value or os.getcwd()))
    except (OSError, TypeError, ValueError):
        return str(value)


# Move, rather than merely add, the expected root to the front. This prevents an
# unrelated earlier UTILITIES package from executing before the location check.
_project_root_key = _path_key(_PROJECT_ROOT)
sys.path[:] = [entry for entry in sys.path if _path_key(entry) != _project_root_key]
sys.path.insert(0, _PROJECT_ROOT)

_EXPECTED_UTILITIES_DIR = _path_key(os.path.join(_PROJECT_ROOT, "UTILITIES"))
_UTILITIES_SPEC = importlib.util.find_spec("UTILITIES")
_UTILITIES_ORIGIN = getattr(_UTILITIES_SPEC, "origin", None)
_ACTUAL_UTILITIES_DIR = (
    _path_key(os.path.dirname(_UTILITIES_ORIGIN)) if _UTILITIES_ORIGIN else ""
)
if _ACTUAL_UTILITIES_DIR != _EXPECTED_UTILITIES_DIR:
    raise ImportError(
        "Refusing to load UTILITIES from an unexpected location: "
        f"{_ACTUAL_UTILITIES_DIR!r}; expected {_EXPECTED_UTILITIES_DIR!r}."
    )

import UTILITIES as _utilities  # noqa: E402
from UTILITIES import (  # noqa: E402
    VerifyMode,
    add_common_portal_args,
    add_config_arg,
    add_email_args,
    add_target_arg,
    build_secure_session,
    build_timestamp,
    connect_gis,
    ensure_output_dir,
    finalize_config,
    gather_attachments,
    get_max_email_mb,
    get_platform,
    provider_for_gis,
    redact_token,
    send_email,
    setup_logging,
    start_timer,
    subject_line,
    timing_block,
    verify_tls_handshake,
)

# Defense in depth after import as well.
if _path_key(os.path.dirname(getattr(_utilities, "__file__", ""))) != \
        _EXPECTED_UTILITIES_DIR:
    raise ImportError(
        "UTILITIES changed location during import; refusing to continue."
    )

try:  # noqa: E402
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Alignment, Font
except ImportError as exc:  # pragma: no cover - deployment dependency
    raise SystemExit(
        "ERROR: openpyxl is required. Install it with: pip install openpyxl"
    ) from exc


REQUIRED_SETTINGS = ("portal-url", "username", "keyring-service", "output-dir")
REQUEST_TIMEOUT = 30
MAX_ROLE_PAGES = 10_000

# Column and row constants for the template.
ROLE_START_COL = 4       # Column D (one-indexed).
PRIVILEGE_START_ROW = 2  # Row 1 is the role-name header.
PRIVILEGE_COL = 3        # Column C contains labels and is never modified.

_TLS_MODE_CHOICES = {
    "standard": VerifyMode.STANDARD_CA,
    "os-trust-store": VerifyMode.OS_TRUST_STORE,
    "custom-ca-bundle": VerifyMode.CUSTOM_CA_BUNDLE,
}

PRIVILEGE_DATA = [
    ("portal:user:useAIAssistants", "Use generative AI and Arcgis assistants"),
    ("portal:user:viewOrgUsers", "View members of the organization"),
    ("portal:user:arcgispro:takeLicenseOffline", "Take Arcgis Pro license offline"),
    ("portal:user:createGroup", "Create, edit, and delete own groups"),
    ("portal:user:joinGroup", "Join groups within the organization"),
    ("portal:user:joinNonOrgGroup", "Join groups external to the organization"),
    ("portal:user:viewOrgGroups", "View groups shared with the organization"),
    ("portal:user:invitePartneredCollaborationMembers", "Invite members from partnered collaboration organizations to groups"),
    ("portal:user:addExternalMembersToGroup", "Create groups that allow members from other organizations"),
    ("portal:user:createItem", "Create, edit, and delete own content"),
    ("portal:publisher:publishFeatures", "Publish hosted feature layers from shapefiles, csvs, etc"),
    ("portal:publisher:publishTiles", "Publish hosted tile layers from tile packages, features, etc"),
    ("portal:publisher:publishScenes", "Publish hosted scene layers"),
    ("portal:publisher:publishTiledImagery", "Publish hosted tiled imagery layers"),
    ("portal:publisher:publishDynamicImagery", "Publish hosted dynamic imagery layers"),
    ("portal:publisher:publishServerServices", "Publish layers powered by services from federated server sites"),
    ("portal:user:viewOrgItems", "View content shared to the organization"),
    ("portal:publisher:bulkPublishFromDataStores", "Publish web layers from a registered data store"),
    ("portal:publisher:registerDataStores", "Register data stores to the portal"),
    ("premium:publisher:createNotebooks", "Create and edit interactive notebooks"),
    ("premium:publisher:scheduleNotebooks", "Schedule future automated runs of a notebook"),
    ("portal:user:viewTracks", "View members' location tracks"),
    ("portal:user:reassignItems", "Reassign own content to another member"),
    ("portal:user:receiveItems", "Receive content that is reassigned by another member"),
    ("portal:publisher:createDataPipelines", "Create, edit, and run data pipelines"),
    ("portal:user:generateApiTokens", "Generate api keys"),
    ("portal:user:assignPrivilegesToApps", "Assign privileges to Oauth 2.0 applications"),
    ("portal:user:createWorkflow", "Create workflow manager items"),
    ("portal:user:shareToGroup", "Share content to groups"),
    ("portal:user:shareToOrg", "Share content to the organization"),
    ("portal:user:shareToPublic", "Share content publicly"),
    ("portal:user:shareGroupToOrg", "Make groups discoverable by the organization"),
    ("portal:user:shareGroupToPublic", "Make groups discoverable by the public"),
    ("opendata:user:designateGroup", "Designate groups as being available for use in open data"),
    ("premium:user:geocode", "Perform large-volume geocoding tasks with the ESRI world geocoder"),
    ("premium:user:networkanalysis", "Perform network analysis tasks such as routing and drive-time areas"),
    ("premium:user:spatialanalysis", "Perform spatial (feature) analysis tasks"),
    ("premium:user:print:customlayout", "Print from layout items"),
    ("premium:user:geoenrichment", "Geoenrich features"),
    ("premium:user:demographics", "Make use of premium demographic data"),
    ("premium:publisher:rasteranalysis", "Perform imagery and raster analysis tasks"),
    ("premium:publisher:createAdvancedNotebooks", "Import and use arcpy modules in notebooks"),
    ("portal:user:runWebTool", "Run web tools"),
    ("premium:user:featurereport", "Create feature reports in Survey123"),
    ("features:user:edit", "Edit features in editable layers that are not public"),
    ("features:user:fullEdit", "Add, delete, and update features in an editable, hosted feature layer"),
    ("features:user:manageVersions", "Manage branch versions accessed through arcgis server feature layers"),
    ("portal:admin:createGPWebhook", "Create, edit, and delete feature layer webhooks"),
    ("portal:admin:viewUsers", "View full member account information within the organization"),
    ("portal:admin:updateUsers", "Reset passwords, update member account information and categories"),
    ("portal:admin:deleteUsers", "Delete member accounts within the organization"),
    ("portal:admin:inviteUsers", "Invite members to the organization"),
    ("portal:admin:disableUsers", "Enable and disable member accounts"),
    ("portal:admin:manageRoles", "Change the role a member account is assigned"),
    ("portal:admin:manageLicenses", "Assign licenses to members of the organization"),
    ("portal:admin:categorizeItems", "Configure organization member categories"),
    ("portal:admin:viewGroups", "View all groups within the organization"),
    ("portal:admin:updateGroups", "Update groups within the organization"),
    ("portal:admin:deleteGroups", "Delete groups within the organization"),
    ("portal:admin:reassignGroups", "Reassign groups to other members within the organization"),
    ("portal:admin:assignToGroups", "Assign, update, and remove members from groups"),
    ("portal:admin:manageEnterpriseGroups", "Link group membership to an organization-specific group"),
    ("portal:admin:createUpdateCapableGroup", "Create groups with update capabilities"),
    ("portal:admin:createLeavingDisallowedGroup", "Create and own groups that do not leave (administrative groups)"),
    ("portal:admin:viewItems", "View all content within the organization"),
    ("portal:admin:updateItems", "Update and categorize content and edit hosted feature layers in the organization"),
    ("portal:admin:deleteItems", "Delete content within the organization"),
    ("portal:admin:reassignItems", "Reassign content to other members within the organization"),
    ("portal:admin:updateItemCategorySchema", "Configure organization content categories"),
    ("portal:publisher:publishServerGPServices", "Publish web tools 1"),
    ("portal:admin:shareToOrg", "Share content owned by other members in the organization with the organization"),
    ("portal:admin:shareToPublic", "Share content owned by other members in the organization with the public"),
    ("portal:admin:createReports", "Create and manage administrative reports for the organization"),
    ("portal:admin:manageWebhooks", "Create, edit, and delete all webhooks in the organization"),
    ("marketplace:admin:manage", "Create listings, list items, manage subscriptions within arcgis marketplace"),
    ("marketplace:admin:purchase", "Send purchase requests and access free products in arcgis marketplace"),
    ("marketplace:admin:startTrial", "Start trial subscriptions within arcgis marketplace"),
    ("portal:admin:manageSecurity", "Manage the organization's security and infrastructure settings"),
    ("portal:admin:manageWebsite", "Manage the organization's website settings"),
    ("portal:admin:manageCollaborations", "Manage the organization's collaborations"),
    ("portal:admin:manageCredits", "Manage the organization's credit budgeting settings"),
    ("portal:admin:changeUserRoles", "Change the role a member account is assigned"),
    ("portal:admin:manageUtilityServices", "Manage the organization's utility service settings"),
]


class RoleFetchError(RuntimeError):
    """A sanitized failure that makes the role inventory incomplete."""

    def __init__(self, message):
        super().__init__(redact_token(str(message)))


class MatrixBuildError(RuntimeError):
    """A structural or write failure while building the workbook."""


class RunFailure(RuntimeError):
    """A controlled fatal failure reported through the normal run email."""

    def __init__(self, message):
        super().__init__(redact_token(str(message)))


def _normalize_text(value):
    return str(value or "").strip().casefold()


def _normalize_https_base(raw, label="Portal URL"):
    """Validate and normalize a credential-free HTTPS Portal/AGOL base URL."""
    text = str(raw or "").strip()
    if not text:
        raise ValueError(f"{label} is blank.")
    if any(ch.isspace() or ord(ch) < 32 for ch in text):
        raise ValueError(f"{label} contains whitespace or a control character.")
    if "\\" in text:
        raise ValueError(f"{label} must not contain a backslash.")
    try:
        parts = urlsplit(text)
    except ValueError as exc:
        raise ValueError(f"{label} could not be parsed: {text!r}") from exc
    if parts.scheme.casefold() != "https":
        raise ValueError(f"{label} must use HTTPS: {text!r}")
    if not parts.hostname:
        raise ValueError(f"{label} contains no hostname: {text!r}")
    try:
        _ = parts.port
    except ValueError as exc:
        raise ValueError(f"{label} contains an invalid port: {text!r}") from exc
    if parts.username is not None or parts.password is not None:
        raise ValueError(f"{label} must not contain user-info credentials.")
    if parts.query or parts.fragment:
        raise ValueError(f"{label} must not contain a query string or fragment.")

    raw_path = parts.path or ""
    lowered_path = raw_path.casefold()
    if "%2f" in lowered_path or "%5c" in lowered_path:
        raise ValueError(f"{label} path must not contain encoded slashes.")
    decoded_path = unquote(raw_path)
    if "\\" in decoded_path or any(ord(ch) < 32 for ch in decoded_path):
        raise ValueError(f"{label} path contains a backslash or control character.")
    decoded_segments = [segment for segment in decoded_path.split("/") if segment]
    if any(segment in {".", ".."} for segment in decoded_segments):
        raise ValueError(f"{label} path contains a traversal component.")

    # Keep the original encoding while normalizing duplicate/trailing slashes.
    raw_segments = [segment for segment in raw_path.split("/") if segment]
    path = "/" + "/".join(raw_segments) if raw_segments else ""
    if path.casefold().endswith("/home"):
        path = path[:-5]
    return urlunsplit(("https", parts.netloc, path, "", "")).rstrip("/")


def _sharing_rest_url(base_url, relative_path):
    base = _normalize_https_base(base_url)
    if base.casefold().endswith("/sharing/rest"):
        prefix = base
    else:
        prefix = base + "/sharing/rest"
    return prefix + "/" + str(relative_path).lstrip("/")


def build_text_to_code(logger=None):
    """Build label -> frozenset(code aliases), retaining the legacy name."""
    text_to_codes = {}
    for esri_json, short_text in PRIVILEGE_DATA:
        code = _normalize_text(esri_json)
        label = _normalize_text(short_text)
        if not code or not label:
            continue
        text_to_codes.setdefault(label, set()).add(code)

    result = {label: frozenset(codes)
              for label, codes in text_to_codes.items()}
    if logger is not None:
        code_count = sum(len(codes) for codes in result.values())
        alias_labels = sum(1 for codes in result.values() if len(codes) > 1)
        logger.info(
            "Built privilege mapping from embedded data: %d label(s), %d "
            "code(s), %d label(s) with version aliases.",
            len(result), code_count, alias_labels,
        )
    return result


def _parse_rest_integer(value, label):
    """Parse an ArcGIS pagination integer without silently truncating values."""
    if isinstance(value, bool):
        raise RoleFetchError(f"{label} is not an integer: {value!r}")
    if isinstance(value, float) and not value.is_integer():
        raise RoleFetchError(f"{label} is not an integer: {value!r}")
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise RoleFetchError(f"{label} is not an integer: {value!r}") from exc


def _format_rest_error(error):
    if not isinstance(error, dict):
        return redact_token(f"REST error: {error}")
    code = error.get("code", "?")
    message = error.get("message", "unknown REST error")
    details = error.get("details") or []
    if isinstance(details, list):
        detail_text = "; ".join(str(value) for value in details if value)
    else:
        detail_text = str(details)
    rendered = f"REST error {code}: {message}"
    if detail_text:
        rendered += f" ({detail_text})"
    return redact_token(rendered)


def _refresh_token_provider(token_provider, gis=None):
    """Best-effort SDK refresh, followed by a mandatory provider refresh."""
    if gis is not None:
        try:
            # This mirrors the project's other direct-REST scanner: give the SDK
            # an authenticated call on which to renew its managed token.
            _ = gis.properties.get("name")
        except Exception:
            pass
    try:
        token_provider.refresh()
        token = token_provider.get()
    except Exception as exc:
        raise RoleFetchError(f"token refresh failed: {exc}") from exc
    if not token:
        raise RoleFetchError("token refresh returned no token")


def _request_role_json(session, url, params, token_provider, *, gis=None):
    """GET one role page, block redirects, and retry one auth failure."""
    for attempt in range(2):
        response = None
        try:
            request_params = dict(params or {})
            request_params.update(token_provider.auth_params())
            response = session.get(
                url,
                params=request_params,
                timeout=REQUEST_TIMEOUT,
                allow_redirects=False,
            )
        except Exception as exc:
            raise RoleFetchError(f"role REST request failed: {exc}") from exc

        try:
            status = int(getattr(response, "status_code", 0) or 0)
            if 300 <= status < 400:
                raise RoleFetchError(
                    f"role REST endpoint returned redirect HTTP {status}; "
                    "token-bearing redirects are refused"
                )
            if status in (401, 403) and attempt == 0:
                _refresh_token_provider(token_provider, gis=gis)
                continue
            if status >= 400:
                reason = getattr(response, "reason", "") or ""
                raise RoleFetchError(
                    f"role REST endpoint returned HTTP {status} {reason}".strip()
                )
            try:
                data = response.json()
            except Exception as exc:
                raise RoleFetchError(
                    f"role REST endpoint returned invalid JSON: {exc}"
                ) from exc
            if not isinstance(data, dict):
                raise RoleFetchError(
                    "role REST endpoint returned a non-object JSON value"
                )

            error = data.get("error")
            if error is not None:
                rest_code = error.get("code") if isinstance(error, dict) else None
                try:
                    rest_code = int(rest_code)
                except (TypeError, ValueError):
                    pass
                if rest_code in (498, 499) and attempt == 0:
                    _refresh_token_provider(token_provider, gis=gis)
                    continue
                raise RoleFetchError(_format_rest_error(error))
            return data
        finally:
            close = getattr(response, "close", None)
            if callable(close):
                try:
                    close()
                except Exception:
                    pass

    raise RoleFetchError("role REST request remained unauthorized after refresh")


def _clean_role(role, page_number, role_index):
    if not isinstance(role, dict):
        raise RoleFetchError(
            f"role page {page_number} entry {role_index} is not an object"
        )
    cleaned = dict(role)
    role_id = str(cleaned.get("id") or "").strip()
    role_name = str(cleaned.get("name") or "").strip()
    if not role_id and not role_name:
        raise RoleFetchError(
            f"role page {page_number} entry {role_index} has no id or name"
        )
    for label, value in (("id", role_id), ("name", role_name)):
        if value and any(ord(ch) < 32 for ch in value):
            raise RoleFetchError(
                f"role page {page_number} entry {role_index} {label} "
                "contains a control character"
            )

    privileges = cleaned.get("privileges")
    if privileges is None:
        privileges = []
    if not isinstance(privileges, list):
        raise RoleFetchError(
            f"role {role_name or role_id!r} returned a non-list privileges value"
        )
    cleaned_privileges = []
    for value in privileges:
        if value is None:
            continue
        rendered = str(value).strip()
        if not rendered:
            continue
        if any(ord(ch) < 32 for ch in rendered):
            raise RoleFetchError(
                f"role {role_name or role_id!r} returned a privilege code "
                "containing a control character"
            )
        cleaned_privileges.append(rendered)

    cleaned["id"] = role_id
    cleaned["name"] = role_name
    cleaned["privileges"] = cleaned_privileges
    return cleaned


def fetch_all_roles(base_url, session, token_provider, logger=None, gis=None):
    """Fetch every role page; any failed/malformed page invalidates the run."""
    roles_url = _sharing_rest_url(
        base_url, "/portals/self/roles"
    )
    common_params = {
        "f": "json",
        "num": 1000,
        "returnPrivileges": "true",
        "all": "true",
    }

    roles = []
    reported_total = None
    start = 1
    seen_starts = set()
    page_number = 0

    while True:
        page_number += 1
        if page_number > MAX_ROLE_PAGES:
            raise RoleFetchError(
                f"role pagination exceeded {MAX_ROLE_PAGES:,} pages"
            )
        if start in seen_starts:
            raise RoleFetchError(
                f"role pagination repeated nextStart={start}"
            )
        seen_starts.add(start)

        params = dict(common_params)
        params["start"] = start
        data = _request_role_json(
            session, roles_url, params, token_provider, gis=gis
        )

        if data.get("total") is not None:
            page_total = _parse_rest_integer(
                data.get("total"), "role endpoint total"
            )
            if page_total < 0:
                raise RoleFetchError(
                    f"role endpoint returned negative total={page_total}"
                )
            if reported_total is None:
                reported_total = page_total
            elif page_total != reported_total:
                raise RoleFetchError(
                    f"role endpoint changed total from {reported_total} to "
                    f"{page_total} during pagination"
                )

        if "roles" not in data:
            if reported_total == 0:
                page_roles = []
            else:
                raise RoleFetchError(
                    f"role page {page_number} returned no 'roles' field"
                )
        else:
            page_roles = data.get("roles")
        if not isinstance(page_roles, list):
            raise RoleFetchError(
                f"role page {page_number} returned a non-list roles value"
            )

        for index, role in enumerate(page_roles, start=1):
            roles.append(_clean_role(role, page_number, index))

        raw_next = data.get("nextStart", -1)
        if raw_next is None:
            raw_next = -1
        next_start = _parse_rest_integer(
            raw_next, f"role page {page_number} nextStart"
        )

        if logger is not None:
            logger.debug(
                "Fetched role page %d: start=%d, rows=%d, nextStart=%d.",
                page_number, start, len(page_roles), next_start,
            )

        if next_start == -1:
            break
        if next_start <= 0 or next_start == start:
            raise RoleFetchError(
                f"role page {page_number} returned non-progressing "
                f"nextStart={next_start}"
            )
        if not page_roles:
            raise RoleFetchError(
                f"role page {page_number} was empty but requested continuation "
                f"at nextStart={next_start}"
            )
        start = next_start

    if reported_total is not None and len(roles) < reported_total:
        raise RoleFetchError(
            f"role pagination returned {len(roles)} row(s), fewer than the "
            f"reported total of {reported_total}"
        )
    unique_count = len(dedupe_roles(roles))
    if reported_total is not None and unique_count != reported_total:
        raise RoleFetchError(
            f"role pagination produced {unique_count} unique role(s), but the "
            f"endpoint reported a total of {reported_total}"
        )
    if logger is not None:
        logger.info(
            "Fetched %d role row(s) across %d page(s).",
            len(roles), page_number,
        )
    return roles


def fetch_builtin_role_ids(base_url, session, token_provider, logger=None,
                           gis=None):
    """Return exact built-in role ids from portals/self.defaultUserRoles."""
    info_url = _sharing_rest_url(base_url, "/portals/self")
    data = _request_role_json(
        session, info_url, {"f": "json"}, token_provider, gis=gis
    )
    raw_roles = data.get("defaultUserRoles")
    if raw_roles is None:
        raw_roles = []
    if not isinstance(raw_roles, list):
        raise RoleFetchError(
            "portals/self returned a non-list defaultUserRoles value"
        )

    role_ids = set()
    malformed = 0
    for value in raw_roles:
        if not isinstance(value, dict):
            malformed += 1
            continue
        role_id = str(value.get("role") or value.get("id") or "").strip()
        if role_id and not any(ord(ch) < 32 for ch in role_id):
            role_ids.add(role_id.casefold())
        else:
            malformed += 1
    if malformed:
        raise RoleFetchError(
            f"portals/self returned {malformed} malformed default role "
            "record(s)"
        )
    if logger is not None:
        logger.info("Resolved %d exact built-in role id(s).", len(role_ids))
        if not role_ids:
            logger.warning(
                "portals/self supplied no default role ids; built-in filtering "
                "will rely on each role's explicit roleType field."
            )
    return frozenset(role_ids)


def _role_identity(role):
    role_id = str(role.get("id") or "").strip()
    if role_id:
        return ("id", role_id.casefold())
    name = str(role.get("name") or "").strip()
    return ("name", name.casefold()) if name else None


def dedupe_roles(all_roles):
    """Deduplicate in first-seen order and retain every reported privilege."""
    by_identity = {}
    order = []
    for role in all_roles or []:
        if not isinstance(role, dict):
            continue
        key = _role_identity(role)
        if key is None:
            continue
        incoming = dict(role)
        incoming_privileges = [
            str(value).strip() for value in (incoming.get("privileges") or [])
            if value is not None and str(value).strip()
        ]
        incoming["privileges"] = incoming_privileges

        if key not in by_identity:
            by_identity[key] = incoming
            order.append(key)
            continue

        existing = by_identity[key]
        seen = {_normalize_text(value)
                for value in (existing.get("privileges") or [])}
        merged_privileges = list(existing.get("privileges") or [])
        for privilege in incoming_privileges:
            normalized = _normalize_text(privilege)
            if normalized not in seen:
                seen.add(normalized)
                merged_privileges.append(privilege)
        existing["privileges"] = merged_privileges

        # Preserve an explicit built-in classification from either duplicate.
        if _normalize_text(incoming.get("roleType")) == "builtin":
            existing["roleType"] = incoming.get("roleType")
        for field in ("name", "id", "description"):
            if not existing.get(field) and incoming.get(field):
                existing[field] = incoming[field]

    return [by_identity[key] for key in order]


def filter_custom_roles(roles, builtin_role_ids=()):
    """Exclude roles explicitly classified or exactly identified as built in."""
    builtin_ids = {str(value).strip().casefold()
                   for value in (builtin_role_ids or ()) if str(value).strip()}
    return [
        role for role in roles
        if _normalize_text(role.get("roleType")) != "builtin"
        and str(role.get("id") or "").strip().casefold() not in builtin_ids
    ]


def _select_privilege_sheet(workbook):
    if "MATRIX" not in workbook.sheetnames:
        raise MatrixBuildError(
            "ROLES_TEMPLATE.xlsx does not contain the required 'MATRIX' sheet"
        )
    return workbook["MATRIX"]


def _mapping_codes(value):
    if value is None:
        return frozenset()
    if isinstance(value, str):
        normalized = _normalize_text(value)
        return frozenset((normalized,)) if normalized else frozenset()
    try:
        values = value
        return frozenset(
            normalized for normalized in (_normalize_text(item) for item in values)
            if normalized
        )
    except TypeError:
        normalized = _normalize_text(value)
        return frozenset((normalized,)) if normalized else frozenset()


def _matrix_role_entries(roles):
    entries = []
    for role in roles:
        role_id = str(role.get("id") or "").strip()
        role_name = str(role.get("name") or role_id).strip()
        if not role_name:
            raise MatrixBuildError("A role has neither a display name nor an id.")
        if any(ord(ch) < 32 for ch in role_name):
            raise MatrixBuildError(
                f"Role {role_id or '<unknown>'!r} contains a control character."
            )
        raw_privileges = role.get("privileges") or []
        if not isinstance(raw_privileges, (list, tuple, set, frozenset)):
            raise MatrixBuildError(
                f"Role {role_name!r} has a non-list privileges value."
            )
        privileges = frozenset(
            normalized
            for normalized in (
                _normalize_text(value) for value in raw_privileges
            )
            if normalized
        )
        entries.append({
            "id": role_id,
            "name": role_name,
            "privileges": privileges,
        })
    # Python's sort is stable, preserving API order for equal privilege counts.
    return sorted(entries, key=lambda entry: len(entry["privileges"]))


def _clear_existing_matrix_values(worksheet, last_role_col):
    last_col = max(worksheet.max_column, last_role_col, ROLE_START_COL)
    last_row = max(worksheet.max_row, PRIVILEGE_START_ROW)
    for row in range(1, last_row + 1):
        for col in range(ROLE_START_COL, last_col + 1):
            cell = worksheet.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None
            cell.hyperlink = None
            cell.comment = None


def _force_literal_text(cell, value):
    """Write untrusted role names as literal strings, never Excel formulas."""
    cell.value = str(value)
    cell.data_type = "s"


def _save_workbook_atomically(workbook, destination):
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{destination.stem}_",
        suffix=".xlsx",
        dir=str(destination.parent),
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        workbook.save(str(temp_path))
        # Release the source workbook before an in-place os.replace on Windows.
        workbook.close()
        try:
            os.chmod(temp_path, 0o600)
        except OSError:
            pass

        # Re-open once before publication so a truncated ZIP is never installed.
        check = openpyxl.load_workbook(str(temp_path), read_only=True)
        check.close()
        os.replace(str(temp_path), str(destination))
        try:
            os.chmod(destination, 0o600)
        except OSError:
            pass
    finally:
        try:
            temp_path.unlink()
        except FileNotFoundError:
            pass
        except OSError:
            # The destination has already been published (or the primary error
            # is more useful); temp cleanup must not mask either outcome.
            pass


def populate_template_xlsx(template_path, roles, text_to_code, logger=None,
                           output_path=None):
    """Populate MATRIX and publish atomically; Column C is never changed."""
    source = Path(template_path)
    destination = Path(output_path) if output_path is not None else source
    workbook = None
    try:
        if not source.is_file():
            raise MatrixBuildError(f"Template not found: {source}")
        workbook = openpyxl.load_workbook(str(source))
        worksheet = _select_privilege_sheet(workbook)

        template_rows = []
        unmatched_labels = []
        for row in range(PRIVILEGE_START_ROW, worksheet.max_row + 1):
            original_label = worksheet.cell(
                row=row, column=PRIVILEGE_COL
            ).value
            if original_label is None or not str(original_label).strip():
                continue
            label = str(original_label).strip()
            codes = _mapping_codes(text_to_code.get(_normalize_text(label)))
            template_rows.append((row, label, codes))
            if not codes:
                unmatched_labels.append(label)

        if not template_rows:
            raise MatrixBuildError(
                "The MATRIX sheet contains no privilege labels in column C."
            )
        matched = sum(1 for _, _, codes in template_rows if codes)
        if matched == 0:
            raise MatrixBuildError(
                "No MATRIX privilege label matched the embedded mapping."
            )

        entries = _matrix_role_entries(roles)
        if not entries:
            raise MatrixBuildError("No custom roles were supplied for the matrix.")

        last_role_col = ROLE_START_COL + len(entries) - 1
        _clear_existing_matrix_values(worksheet, last_role_col)

        represented_codes = frozenset(
            code for _row, _label, codes in template_rows for code in codes
        )
        unknown_by_role = []
        for entry in entries:
            unknown = sorted(entry["privileges"] - represented_codes)
            if unknown:
                unknown_by_role.append((entry["name"], unknown))

        names = {}
        for entry in entries:
            names.setdefault(entry["name"].casefold(), []).append(entry["id"])
        duplicate_names = {
            name: ids for name, ids in names.items() if len(ids) > 1
        }

        for role_index, entry in enumerate(entries):
            col = ROLE_START_COL + role_index
            cell = worksheet.cell(row=1, column=col)
            _force_literal_text(cell, entry["name"])
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
                textRotation=90,
            )
            col_letter = openpyxl.utils.get_column_letter(col)
            worksheet.column_dimensions[col_letter].width = 4
        worksheet.row_dimensions[1].height = 150

        for template_row, _label, mapped_codes in template_rows:
            if not mapped_codes:
                continue
            for role_index, entry in enumerate(entries):
                if mapped_codes.intersection(entry["privileges"]):
                    cell = worksheet.cell(
                        row=template_row,
                        column=ROLE_START_COL + role_index,
                    )
                    cell.value = "✓"
                    cell.alignment = Alignment(
                        horizontal="center", vertical="center"
                    )

        if logger is not None:
            logger.info(
                "Read %d privilege row(s) from template column C; %d matched "
                "the embedded code mapping.",
                len(template_rows), matched,
            )
            if unmatched_labels:
                logger.warning(
                    "%d template privilege label(s) have no embedded code and "
                    "were left blank: %s",
                    len(unmatched_labels), "; ".join(unmatched_labels[:25]),
                )
            if unknown_by_role:
                unknown_codes = sorted({
                    code for _name, codes in unknown_by_role for code in codes
                })
                logger.warning(
                    "%d role(s) contain %d privilege code(s) not represented "
                    "by the current template mapping; those codes cannot be "
                    "shown without changing column C: %s",
                    len(unknown_by_role), len(unknown_codes),
                    ", ".join(unknown_codes[:50]),
                )
            if duplicate_names:
                logger.warning(
                    "%d duplicate role display name(s) were retained as "
                    "separate columns using role ids internally.",
                    len(duplicate_names),
                )
            logger.info(
                "Roles ordered by permission count: %s",
                " -> ".join(
                    f"{entry['name']} ({len(entry['privileges'])} privs)"
                    for entry in entries
                ),
            )

        _save_workbook_atomically(workbook, destination)
        if logger is not None:
            logger.info(
                "Created %s with checkmarks for %d role(s) across %d "
                "privilege row(s). Column C was left unchanged.",
                destination, len(entries), len(template_rows),
            )
        return True
    except Exception as exc:
        if logger is not None:
            logger.error("Error populating template: %s", redact_token(str(exc)))
        return False
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "Generate a privilege matrix for custom Portal/AGOL roles in "
            "ROLES_TEMPLATE.xlsx. Direct REST and ArcGIS SDK TLS verification "
            "both follow --tls-mode."
        )
    )
    add_config_arg(parser)
    add_target_arg(parser)
    add_common_portal_args(parser)
    add_email_args(parser)
    # The shared helper's help text says "CSV"; this script writes XLSX. The
    # argument shape remains identical so finalize_config() resolves it normally.
    parser.add_argument(
        "--output-dir",
        default=None,
        help=(
            "Directory where the timestamped Excel matrix and log are written. "
            "If omitted, taken from the --config org toml ('output-dir')."
        ),
    )

    tls_group = parser.add_argument_group("TLS verification")
    tls_group.add_argument(
        "--tls-mode",
        choices=sorted(_TLS_MODE_CHOICES),
        default="standard",
        help=(
            "Verification for both direct role REST calls and connect_gis(): "
            "'standard' (certifi), 'os-trust-store' (requires truststore), or "
            "'custom-ca-bundle' (requires --ca-bundle-path)."
        ),
    )
    tls_group.add_argument(
        "--ca-bundle-path",
        default=None,
        help=(
            "PEM CA bundle required by --tls-mode custom-ca-bundle; ignored "
            "for the other modes."
        ),
    )
    return parser.parse_args()


def _reserve_run_paths(output_dir, org_prefix, platform, target, timestamp):
    for sequence in range(1000):
        run_stamp = timestamp if sequence == 0 else f"{timestamp}_{sequence:02d}"
        log_path = Path(output_dir) / (
            f"{org_prefix}_{platform}_PORTAL_ROLES_{run_stamp}.log"
        )
        output_path = Path(output_dir) / (
            f"{org_prefix}_ROLES_{target}_{run_stamp}.xlsx"
        )
        if output_path.exists():
            continue
        try:
            fd = os.open(
                str(log_path), os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600
            )
        except FileExistsError:
            continue
        except OSError as exc:
            raise RunFailure(
                f"Could not reserve log file {log_path}: {exc}"
            ) from exc
        try:
            os.close(fd)
        except OSError as exc:
            raise RunFailure(
                f"Could not close reserved log file {log_path}: {exc}"
            ) from exc
        return run_stamp, str(log_path), output_path
    raise RunFailure(
        "Could not reserve unique log/output names after 1,000 attempts."
    )


def _effective_ca_path(args, mode):
    return args.ca_bundle_path if mode is VerifyMode.CUSTOM_CA_BUNDLE else None


def _build_direct_session(args, mode, logger):
    return build_secure_session(
        mode,
        ca_bundle_path=_effective_ca_path(args, mode),
        pool_connections=16,
        pool_maxsize=16,
        retry_total=2,
        logger=logger,
    )


def _connect_verified_gis(args, mode, logger):
    return connect_gis(
        args.portal_url,
        args.username,
        args.keyring_service,
        logger,
        verify_mode=mode,
        ca_bundle_path=_effective_ca_path(args, mode),
    )


def _gis_base_url(gis, fallback):
    value = getattr(gis, "url", None)
    if not value:
        portal = getattr(gis, "_portal", None)
        value = getattr(portal, "url", None)
    return _normalize_https_base(value or fallback, "Connected GIS URL")


def _send_abort_email(args, logger, timestamp, platform, log_file_path,
                      start_wallclock, start_perf, reason):
    if not args.email_to:
        return
    safe_reason = redact_token(str(reason))
    body = (
        "Portal Roles Privilege Matrix\r\n"
        f"Organization: {args.org_name}\r\n"
        f"Platform: {platform}\r\n"
        f"Run timestamp: {timestamp}\r\n"
        "Result: ABORTED\r\n\r\n"
        f"Reason:\r\n{safe_reason}\r\n\r\n"
        f"{timing_block(start_wallclock, start_perf)}"
    )
    try:
        send_email(
            subject_line(
                "Portal Roles Matrix", timestamp, platform,
                org_name=args.org_name,
            ),
            body,
            [log_file_path],
            args.smtp_server,
            args.email_to,
            args.email_from,
            logger,
        )
    except Exception as exc:
        logger.warning("Could not send abort notification email: %s", exc)


def _send_success_email(args, logger, timestamp, platform, log_file_path,
                        output_path, role_count, mapping_count, max_email_mb,
                        start_wallclock, start_perf):
    if not args.email_to:
        return
    attachments, note = gather_attachments(
        log_file_path,
        [str(output_path)],
        args.output_dir,
        logger,
        cap_mb=max_email_mb,
        label="Excel Matrix",
    )
    subject = subject_line(
        "Portal Roles Privilege Matrix",
        timestamp,
        platform=platform,
        org_name=args.org_name,
    )
    body = (
        "Portal Roles Privilege Matrix Report\r\n"
        f"Organization: {args.org_name}\r\n"
        f"Platform: {platform}\r\n"
        f"TLS mode: {args.tls_mode}\r\n"
        f"Run timestamp: {timestamp}\r\n"
        f"Username: {args.username}\r\n\r\n"
        f"Roles processed: {role_count}\r\n"
        f"Privilege labels mapped: {mapping_count}\r\n"
        f"Output file: {output_path.name}\r\n\r\n"
        "The attached Excel file contains a privilege matrix with:\r\n"
        "  - Columns (starting D): All custom roles\r\n"
        "  - Rows: Privileges already defined in the template's column C\r\n"
        "  - Cells: checkmark if the role has that privilege, else empty\r\n\r\n"
        "This can be used for documentation, auditing, and role design.\r\n\r\n"
        f"{timing_block(start_wallclock, start_perf)}\r\n"
        f"{note}"
    )
    send_email(
        subject, body, attachments, args.smtp_server,
        args.email_to, args.email_from, logger,
    )


def _log_start(logger, args, platform, timestamp, log_file_path,
               output_path, max_email_mb):
    logger.info("=" * 70)
    logger.info("PORTAL_ROLES.py - Excel Matrix Builder")
    logger.info("=" * 70)
    logger.info("Organization: %s", args.org_name)
    logger.info("Org config: %s", args.config)
    logger.info("Config target: %s", args.target)
    logger.info("Platform: %s", platform)
    logger.info("Portal URL: %s", args.portal_url)
    logger.info("Username: %s", args.username)
    logger.info("Keyring service: %s", args.keyring_service)
    logger.info("TLS mode: %s", args.tls_mode)
    if args.tls_mode == "custom-ca-bundle":
        logger.info("CA bundle path: %s", args.ca_bundle_path)
    logger.info("Output directory: %s", args.output_dir)
    logger.info("Output workbook: %s", output_path)
    logger.info("Log file: %s", log_file_path)
    logger.info("SMTP server: %s", args.smtp_server)
    logger.info("Email to: %s | from: %s", args.email_to, args.email_from)
    logger.info("Max email attachment size (MB): %s", max_email_mb)
    logger.info("Run timestamp: %s", timestamp)
    logger.info("=" * 70)


def _run(args, logger, timestamp, platform, log_file_path, output_path,
         max_email_mb, start_wallclock, start_perf):
    template_source = Path(_SCRIPT_DIR) / "ROLES_TEMPLATE.xlsx"
    if not template_source.is_file():
        raise RunFailure(
            f"ROLES_TEMPLATE.xlsx not found at {template_source}"
        )

    text_to_code = build_text_to_code(logger)
    mode = _TLS_MODE_CHOICES[args.tls_mode]
    session = None
    try:
        try:
            session = _build_direct_session(args, mode, logger)
        except Exception as exc:
            raise RunFailure(
                f"Could not establish TLS verification strategy "
                f"{args.tls_mode!r}: {exc}"
            ) from exc

        handshake_ok, handshake_detail = verify_tls_handshake(
            args.portal_url, session, logger=logger
        )
        if not handshake_ok:
            raise RunFailure(
                f"TLS handshake verification to {args.portal_url} failed: "
                f"{handshake_detail.get('error')}"
            )

        try:
            gis = _connect_verified_gis(args, mode, logger)
        except Exception as exc:
            raise RunFailure(f"Could not connect to GIS: {exc}") from exc

        token_provider = provider_for_gis(gis, platform)
        logger.info(
            "Token acquired via %s (token present: %s).",
            type(token_provider).__name__, token_provider.is_present(),
        )
        base_url = _gis_base_url(gis, args.portal_url)

        logger.info("STEP 1/2: Fetching all custom roles...")
        try:
            fetched_roles = fetch_all_roles(
                base_url, session, token_provider, logger=logger, gis=gis
            )
        except RoleFetchError as exc:
            raise RunFailure(f"Role inventory is incomplete: {exc}") from exc

        try:
            builtin_role_ids = fetch_builtin_role_ids(
                base_url, session, token_provider, logger=logger, gis=gis
            )
        except RoleFetchError as exc:
            raise RunFailure(
                f"Built-in role identification is incomplete: {exc}"
            ) from exc

        custom_roles = filter_custom_roles(
            dedupe_roles(fetched_roles), builtin_role_ids
        )
        if not custom_roles:
            raise RunFailure("No custom roles were returned by the portal.")
        logger.info("Custom roles to process: %d", len(custom_roles))

        logger.info("STEP 2/2: Populating %s...", output_path)
        if not populate_template_xlsx(
            str(template_source),
            custom_roles,
            text_to_code,
            logger=logger,
            output_path=str(output_path),
        ):
            raise RunFailure("Failed to populate the Excel template.")

        # Write the final run summary before gathering attachments so the log
        # copy sent by email is complete, not missing the success summary.
        logger.info("=" * 70)
        logger.info("RUN SUMMARY")
        logger.info("Organization: %s", args.org_name)
        logger.info("Platform: %s", platform)
        logger.info("TLS mode: %s", args.tls_mode)
        logger.info("Custom roles processed: %d", len(custom_roles))
        logger.info("Privilege labels mapped: %d", len(text_to_code))
        logger.info("Output file: %s", output_path)
        logger.info("Log file: %s", log_file_path)
        logger.info("=" * 70)
        logger.info("Success")

        try:
            _send_success_email(
                args, logger, timestamp, platform, log_file_path,
                output_path, len(custom_roles), len(text_to_code),
                max_email_mb, start_wallclock, start_perf,
            )
        except Exception as exc:
            logger.warning("Could not prepare/send completion email: %s", exc)
        return 0
    finally:
        if session is not None:
            try:
                session.close()
            except Exception as exc:
                logger.warning("Could not close direct REST session: %s", exc)


def main():
    args = parse_args()
    start_wallclock, start_perf = start_timer()
    timestamp = build_timestamp()

    finalize_config(args, required=REQUIRED_SETTINGS)
    try:
        args.portal_url = _normalize_https_base(args.portal_url)
    except ValueError as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(2) from exc

    mode = _TLS_MODE_CHOICES[args.tls_mode]
    if mode is VerifyMode.CUSTOM_CA_BUNDLE:
        if not args.ca_bundle_path:
            print(
                "--ca-bundle-path is required when --tls-mode is "
                "custom-ca-bundle.",
                file=sys.stderr,
            )
            raise SystemExit(2)
        if not Path(args.ca_bundle_path).is_file():
            print(
                f"--ca-bundle-path does not name a file: "
                f"{args.ca_bundle_path}",
                file=sys.stderr,
            )
            raise SystemExit(2)

    platform = get_platform(args.portal_url)
    ensure_output_dir(args.output_dir)
    try:
        timestamp, log_file_path, output_path = _reserve_run_paths(
            args.output_dir,
            args.org_name_safe,
            platform,
            args.target,
            timestamp,
        )
    except RunFailure as exc:
        print(str(exc), file=sys.stderr)
        raise SystemExit(1) from exc

    logger = setup_logging(log_file_path, "portal_roles")
    max_email_mb = get_max_email_mb(args.target, args.config)
    _log_start(
        logger, args, platform, timestamp, log_file_path,
        output_path, max_email_mb,
    )

    try:
        exit_code = _run(
            args, logger, timestamp, platform, log_file_path, output_path,
            max_email_mb, start_wallclock, start_perf,
        )
    except RunFailure as exc:
        safe_reason = redact_token(str(exc))
        logger.error("Run aborted: %s", safe_reason)
        _send_abort_email(
            args, logger, timestamp, platform, log_file_path,
            start_wallclock, start_perf, safe_reason,
        )
        exit_code = 1
    except Exception as exc:
        safe_reason = redact_token(str(exc))
        logger.exception("Unexpected fatal error: %s", safe_reason)
        _send_abort_email(
            args, logger, timestamp, platform, log_file_path,
            start_wallclock, start_perf,
            f"Unexpected failure: {safe_reason}",
        )
        exit_code = 1

    raise SystemExit(exit_code)


if __name__ == "__main__":
    main()
